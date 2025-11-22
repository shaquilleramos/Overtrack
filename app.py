#app.py
from flask import Flask, render_template, request, session, send_file, redirect, url_for
import pandas as pd
import os
import io
from openpyxl.styles import PatternFill, Alignment, Font
from openpyxl.worksheet.table import Table, TableStyleInfo

from utils.procesamiento import procesar_registros, HORARIOS_SEDES

app = Flask(__name__)
app.secret_key = "clave-ultra-secreta"

# Memoria global (sin parquet)
MEMORY = {}

UPLOAD_FOLDER = "uploads"
os.makedirs(UPLOAD_FOLDER, exist_ok=True)


@app.route("/")
def index():
    return render_template("index.html", sedes=list(HORARIOS_SEDES.keys()))

@app.route("/horarios")
def horaios():
    return render_template("horarios.html", sedes=list(HORARIOS_SEDES.keys()))

@app.route("/subir", methods=["POST"])
def subir():
    archivo = request.files.get("archivo_csv")
    sede = request.form.get("sede")

    if not archivo or not archivo.filename:
        return render_template("index.html", mensaje_error="No se seleccionó archivo",
                               sedes=list(HORARIOS_SEDES.keys()))

    if sede not in HORARIOS_SEDES:
        return render_template("index.html", mensaje_error="Seleccione una sede válida",
                               sedes=list(HORARIOS_SEDES.keys()))

    ruta = os.path.join(UPLOAD_FOLDER, archivo.filename)
    archivo.save(ruta)
   

    extension = archivo.filename.lower().split(".")[-1]

    try:
        if extension in ["xlsx", "xls"]:
            df = pd.read_excel(ruta)
        elif extension == "csv":
            try:
                # Intentar primero UTF-8 y detectar separador automáticamente
                df = pd.read_csv(
                    ruta,
                    encoding="utf-8",
                    sep=None,        # autodetecta: , ; | tab
                    engine="python",
                    on_bad_lines="skip"
                )
            except:
                # Si falla, intentar Latin-1 (muy común en Windows)
                df = pd.read_csv(
                    ruta,
                    encoding="latin-1",
                    sep=None,        # autodetecta
                    engine="python",
                    on_bad_lines="skip"
                )

        else:
            return render_template(
                "index.html",
                mensaje_error="Formato no soportado. Use CSV o Excel.",
                sedes=list(HORARIOS_SEDES.keys()),
            )
    except Exception as e:
        return render_template(
            "index.html",
            mensaje_error=f"Error leyendo el archivo: {str(e)}",
            sedes=list(HORARIOS_SEDES.keys()),
        )

    procesado = procesar_registros(df, sede)

    # Guardamos el DF directamente en memoria
    MEMORY["df"] = procesado

    return redirect(url_for("vista_previa"))


@app.route("/vista_previa", methods=["GET", "POST"])
def vista_previa():
    df = MEMORY.get("df")
    if df is None:
        return "No hay datos cargados"

    # Lista de nombres
    nombres = sorted([
    n for n in df["Nombre"].unique()
    if "TOTAL" not in n.upper()
    ])

    # Si no filtra, mostrar todo tal cual viene (ya tenía totales insertados)
    if request.method == "GET":
        return render_template(
            "vista_previa.html",
            tabla=df.to_dict(orient="records"),
            nombres=nombres,
            empleado_seleccionado=None
        )

    # POST: filtrar
    empleado = request.form.get("empleado")

    # --- Si quiere todos, mostrar todo el DF con totales ---
    if not empleado or empleado == "":
        filtrado = df.copy()

    else:
        # Filtrar solo esa persona
        filtrado = df[df["Nombre"] == empleado].copy()

        # Obtener solo sus filas normales
        filas_normales = filtrado[~filtrado["Nombre"].str.contains("TOTAL")].copy()

        # Calcular nuevamente su total de extras
        def to_minutes(x):
            if isinstance(x, str) and "h" in x:
                h, m = x.replace("h", "").replace("m", "").split()
                return int(h)*60 + int(m)
            return 0

        total_min = filas_normales["Horas extras"].apply(to_minutes).sum()
        total_str = f"{total_min//60:02d}h {total_min%60:02d}m"

        # Crear fila TOTAL
        fila_total = {
            "Nombre": f"TOTAL HORAS EXTRAS ({empleado})",
            "Fecha": "",
            "Día": "",
            "Entrada": "",
            "Salida": "",
            "Horas trabajadas": "",
            "Tardanza": "",
            "Horas extras": total_str
        }

        # Volver a unir todo
        filtrado = filas_normales.to_dict("records")
        filtrado.append(fila_total)

        return render_template(
            "vista_previa.html",
            tabla=filtrado,
            nombres=nombres,
            empleado_seleccionado=empleado
        )

    # Render para “todos”
    return render_template(
        "vista_previa.html",
        tabla=filtrado.to_dict(orient="records"),
        nombres=nombres,
        empleado_seleccionado=empleado
    )



@app.route('/descargar_extras')
def descargar_extras():
    nombre = request.args.get("nombre", "todos")

    df = MEMORY.get("df")
    if df is None or df.empty:
        return "No hay datos cargados para exportar", 400

    df = df.copy()

    if nombre != "todos":
        df = df[df["Nombre"] == nombre]

    columnas = [
        "Nombre", "Fecha", "Día", "Entrada", "Salida",
        "Horas trabajadas", "Tardanza", "Horas extras"
    ]

    df = df[columnas]

    # ───────────────────────────────────────────────
    # Convertir "01h 32m" → minutos
    # ───────────────────────────────────────────────
    def parse_horas_extras(valor):
        if isinstance(valor, str) and "h" in valor:
            h, m = valor.replace("h", "").replace("m", "").split()
            return int(h) * 60 + int(m)
        return 0

    df["extra_min"] = df["Horas extras"].apply(parse_horas_extras)
    total_min = df["extra_min"].sum()
    total_horas = f"{total_min // 60:02d}h {total_min % 60:02d}m"

    # Fila total
    fila_total = {
        "Nombre": "",
        "Fecha": "",
        "Día": "",
        "Entrada": "",
        "Salida": "",
        "Horas trabajadas": "",
        "Tardanza": "TOTAL",
        "Horas extras": total_horas,
        "extra_min": total_min
    }

    df_total = pd.concat([df, pd.DataFrame([fila_total])], ignore_index=True)

    df_total = df_total.drop(columns=["extra_min"])

    # ───────────────────────────────────────────────
    # Exportar Excel con estilo
    # ───────────────────────────────────────────────
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df_total.to_excel(writer, index=False, sheet_name="Horas Extras")

        libro = writer.book
        hoja = writer.sheets["Horas Extras"]

        # -------------------------
        # 1️⃣ Crear tabla con estilo
        # -------------------------
        last_row = len(df_total) + 1
        last_col = len(df_total.columns)

        ref = f"A1:{chr(64+last_col)}{last_row}"
        tabla = Table(displayName="TablaExtras", ref=ref)

        estilo = TableStyleInfo(
            name="TableStyleMedium9",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False
        )

        tabla.tableStyleInfo = estilo
        hoja.add_table(tabla)

        # -------------------------
        # 2️⃣ Ajustar ancho de columnas automáticamente
        # -------------------------
        for col in hoja.columns:
            max_length = 0
            col_letter = col[0].column_letter
            for cell in col:
                try:
                    contenido = str(cell.value)
                    max_length = max(max_length, len(contenido))
                except:
                    pass
            hoja.column_dimensions[col_letter].width = max_length + 3

        # -------------------------
        # 3️⃣ Fijar altura de fila
        # -------------------------
        for i in range(1, last_row + 1):
            hoja.row_dimensions[i].height = 22

        # -------------------------
        # 4️⃣ Colorear última fila (TOTAL)
        # -------------------------
        fill_total = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")
        bold_font = Font(bold=True)

        for cell in hoja[last_row]:
            cell.fill = fill_total
            cell.font = bold_font
            cell.alignment = Alignment(horizontal="center")

    output.seek(0)

    return send_file(
        output,
        download_name=f"horas_extras_{nombre}.xlsx",
        as_attachment=True
    )



@app.route('/descargar_llegadas')
def descargar_llegadas():
    nombre = request.args.get("nombre", "todos")

    df = MEMORY.get("df")  # tu DataFrame procesado
    
    if df is None or df.empty:
        return "No hay datos cargados para exportar", 400
    
    df = df.copy()

    # Filtrar por persona
    if nombre != "todos":
        df = df[df["Nombre"] == nombre]

    # Solo estas columnas
    columnas = ["Nombre", "Fecha", "Día", "Entrada"]
    df = df[columnas]

    # Creamos un Excel con estilo
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name="Llegadas")

        libro = writer.book
        hoja = writer.sheets["Llegadas"]

        # ============================
        # 1️⃣ Crear tabla con estilo
        # ============================
        last_row = len(df) + 1
        last_col = len(df.columns)

        ref = f"A1:{chr(64 + last_col)}{last_row}"
        tabla = Table(displayName="TablaLlegadas", ref=ref)

        estilo = TableStyleInfo(
            name="TableStyleMedium9",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False
        )

        tabla.tableStyleInfo = estilo
        hoja.add_table(tabla)

        # ============================
        # 2️⃣ Auto-ajuste de columnas
        # ============================
        for col in hoja.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                try:
                    val = str(cell.value)
                    max_len = max(max_len, len(val))
                except:
                    pass
            hoja.column_dimensions[col_letter].width = max_len + 3

        # ============================
        # 3️⃣ Altura fija de filas
        # ============================
        for i in range(1, last_row + 1):
            hoja.row_dimensions[i].height = 22

    output.seek(0)

    return send_file(
        output,
        download_name=f"llegadas_{nombre}.xlsx",
        as_attachment=True
    )


if __name__ == "__main__":
    app.run(debug=True)
